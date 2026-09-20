#!/usr/bin/env python3
"""Extract registered delegates (rows with email) from SEAMUN allocation matrix v3 layout."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

SKIP_SHEETS = frozenset({"MAIN", "OVERVIEW", "BEG", "INT", "ADV"})
HEADER_LABELS = frozenset({"Country", "Delegation", "Person", "Agency"})


def parse_matrix(path: Path) -> list[dict[str, str]]:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    delegates: list[dict[str, str]] = []
    for sheet in wb.sheetnames:
        if sheet in SKIP_SHEETS:
            continue
        ws = wb[sheet]
        hdr = None
        for ri, row in enumerate(ws.iter_rows(values_only=True), start=1):
            a = (row[0] or "").strip() if row and row[0] else ""
            if a in HEADER_LABELS:
                hdr = ri
                break
        if hdr is None:
            continue
        for row in ws.iter_rows(min_row=hdr + 1, values_only=True):
            delegation = (str(row[0]).strip() if row[0] else "")
            if not delegation or delegation.lower().startswith("total"):
                break
            email = (str(row[7]).strip().lower() if len(row) > 7 and row[7] else "")
            if not email or "@" not in email:
                continue
            delegates.append(
                {
                    "sheet": sheet,
                    "delegation": delegation,
                    "name": (str(row[5]).strip() if len(row) > 5 and row[5] else ""),
                    "email": email,
                    "school": (str(row[6]).strip() if len(row) > 6 and row[6] else ""),
                    "placardCode": (str(row[2]).strip() if len(row) > 2 and row[2] else ""),
                    "status": (str(row[1]).strip() if len(row) > 1 and row[1] else ""),
                }
            )
    return delegates


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("xlsx", type=Path)
    args = parser.parse_args()
    if not args.xlsx.is_file():
        print(f"File not found: {args.xlsx}", file=sys.stderr)
        return 1
    json.dump(parse_matrix(args.xlsx), sys.stdout, indent=2)
    sys.stdout.write("\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
